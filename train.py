from __future__ import division
from __future__ import print_function

import csv
import os
import platform

import numpy as np
import tensorflow.compat.v1 as tf
from openpyxl.workbook import Workbook
from sklearn import metrics

tf.compat.v1.disable_eager_execution()
from utils import *
from models import KGNN
import time

#Set random seed
seed = 2
batch_size=2048
print("seed: ",seed)
np.random.seed(seed)
tf.set_random_seed(seed)

# Settings
dataset = "Drugs"
embedding = 'node2vec'
print("dataset: " + dataset + " | embedding: " + embedding)

flags = tf.app.flags
FLAGS = flags.FLAGS
flags.DEFINE_string('dataset', dataset, 'Dataset string.')  # 'mr','ohsumed','R8','R52'
flags.DEFINE_string('model', 'ConKGNN', 'Model string.')
flags.DEFINE_float('learning_rate', 0.0045, 'Initial learning rate.')
flags.DEFINE_integer('epochs', 200, 'Number of epochs to train.')
flags.DEFINE_integer('batch_size', 2048, 'Size of batches per epoch.')
flags.DEFINE_integer('input_dim', 300, 'Dimension of input.')
flags.DEFINE_integer('hidden', 96, 'Number of units in hidden layer.')  # 32, 64, 96, 128
flags.DEFINE_integer('steps', 2, 'Number of graph layers.')
flags.DEFINE_float('dropout', 0.5, 'Dropout rate (1 - keep probability).')
flags.DEFINE_float('weight_decay', 0, 'Weight for L2 loss on embedding matrix.')  # 5e-4
flags.DEFINE_integer('early_stopping', -1, 'Tolerance for early stopping (# of epochs).')
flags.DEFINE_integer('max_degree', 3, 'Maximum Chebyshev polynomial degree.')  # Not used
path = os.getcwd() + '/data/' + dataset+'/'+embedding+'_knowledge_'

# Load data
train_adj, train_feature, train_y,train_b_mask, val_adj, val_feature,val_y,val_b_mask,test_adj, test_feature, test_y,test_b_mask, = load_data_augment(
    FLAGS.dataset, path)

# Some preprocessing
print('loading training set')
train_adj, train_mask,train_b_mask ,train_b_mask_ud= preprocess_adj_mask(train_adj,train_b_mask)
train_feature = preprocess_features(train_feature)
print('loading validation set')
val_adj, val_mask,val_b_mask, val_b_mask_ud= preprocess_adj_mask(val_adj,val_b_mask)
val_feature = preprocess_features(val_feature)
print('loading test set')
test_adj, test_mask,test_b_mask,test_b_mask_ud = preprocess_adj_mask(test_adj,test_b_mask)
test_feature = preprocess_features(test_feature)

if FLAGS.model == 'ConKGNN':
    model_func = KGNN
else:
    raise ValueError('Invalid argument for model: ' + str(FLAGS.model))

# Define placeholders
placeholders = {
    'support': tf.placeholder(tf.float32, shape=(None, None, None)),
    'features': tf.placeholder(tf.float32, shape=(None, None, FLAGS.input_dim)),
    'mask': tf.placeholder(tf.float32, shape=(None, None, 1)),
    'labels': tf.placeholder(tf.float32, shape=(None, train_y.shape[1])),
    'b_mask': tf.placeholder(tf.float32, shape=(None, None, 1)),
    'b_mask_ud': tf.placeholder(tf.float32, shape=(None, None, 1)),
    'dropout': tf.placeholder_with_default(0., shape=()),
    'num_features_nonzero': tf.placeholder(tf.int32),  # helper variable for sparse dropout
    'supcon_mask': tf.placeholder(tf.float32, shape=(None, None)),
    'supcon_logits_mask': tf.placeholder(tf.float32, shape=(None, None)),
}

# Create model
model = model_func(placeholders, input_dim=FLAGS.input_dim, logging=True)

# Initialize session
sess = tf.Session()

# Define model evaluation function
def evaluate(features, support, mask, labels,background_mask,background_mask_ud,supcon_mask,supcon_logits_mask,placeholders):
    feed_dict_val = construct_feed_dict(features, support, mask, labels,background_mask,background_mask_ud,supcon_mask,supcon_logits_mask, placeholders)
    outs_val = sess.run([model.loss, model.accuracy, model.preds, model.labels, model.outputs,model.y_pres],
                        feed_dict=feed_dict_val)
    return outs_val[0], outs_val[1], outs_val[2], outs_val[3]


# Init variables
sess.run(tf.global_variables_initializer())

cost_val = []
best_val = 0
best_epoch = 0
best_acc = 0
best_cost = 0
test_doc_embeddings = None
preds = None
labels = None
best_message = []
best_outputs=None

print('train start...')
# Train model
for epoch in range(FLAGS.epochs):
    t = time.time()

    # Training step
    indices = np.arange(0, len(train_y))
    np.random.shuffle(indices)
    augment_indices=np.append(indices,indices+len(train_y))

    train_loss, train_acc = 0, 0
    for start in range(0, len(train_y), FLAGS.batch_size):
        end = start + FLAGS.batch_size
        idx = indices[start:end]

        start_augment = start + len(train_y)
        end_augment = start_augment + FLAGS.batch_size
        idx_augment = augment_indices[start_augment:end_augment]

        # Construct feed dictionary
        supcon_mask, supcon_logits_mask=construct_supcon_para(train_y[idx],len(idx))
        feed_dict = construct_feed_dict(np.append(train_feature[idx],train_feature[idx_augment],axis=0), \
                                          np.append(train_adj[idx],train_adj[idx_augment],axis=0), \
                                          np.append(train_mask[idx],train_mask[idx_augment],axis=0), \
                                          train_y[idx],np.append(train_b_mask[idx],train_b_mask[idx_augment],\
                                          axis=0),np.append(train_b_mask_ud[idx],train_b_mask_ud[idx_augment],axis=0),
                                          supcon_mask,supcon_logits_mask,\
                                          placeholders)
        feed_dict.update({placeholders['dropout']: FLAGS.dropout})
        outs = sess.run([model.opt_op, model.loss, model.accuracy], feed_dict=feed_dict)
        train_loss += outs[1] * len(idx)
        train_acc += outs[2] * len(idx)
    train_loss /= len(train_y)
    train_acc /= len(train_y)

    # Validation
    supcon_mask, supcon_logits_mask = construct_supcon_para(val_y, len(val_y))
    val_cost, val_acc, _, _,= evaluate(val_feature, val_adj, val_mask, val_y,val_b_mask,val_b_mask_ud,supcon_mask,supcon_logits_mask,placeholders)
    cost_val.append(val_cost)

    # Test
    supcon_mask, supcon_logits_mask = construct_supcon_para(test_y, len(test_y))
    test_cost, test_acc, pred, labels= evaluate(test_feature, test_adj,test_mask, test_y,test_b_mask,test_b_mask_ud,supcon_mask,supcon_logits_mask,
                                                                                             placeholders)

    if val_acc > best_val:
        best_val = val_acc
        best_epoch = epoch
        best_acc = test_acc
        best_cost = test_cost
        preds = pred
        best_labels = labels

    # Print results
    print("Epoch:", '%04d' % (epoch ), "train_loss=", "{:.5f}".format(train_loss),
          "train_acc=", "{:.5f}".format(train_acc), "val_loss=", "{:.5f}".format(val_cost),
          "val_acc=", "{:.5f}".format(val_acc), "test_acc=", "{:.5f}".format(test_acc),"test_cost=", "{:.5f}".format(test_cost),
          "time=", "{:.5f}".format(time.time() - t))

    if FLAGS.early_stopping > 0 and epoch > FLAGS.early_stopping and cost_val[-1] > np.mean(
            cost_val[-(FLAGS.early_stopping + 1):-1]):
        print("Early stopping...")
        break

print("Optimization Finished!")
print('Best epoch:', best_epoch)
print("Test set results:", "cost=", "{:.5f}".format(best_cost),
      "accuracy=", "{:.5f}".format(best_acc))

print("Test Precision, Recall and F1-Score...")
report1=metrics.classification_report(best_labels, preds, digits=4)
print(report1)
report2=metrics.classification_report(best_labels, preds, digits=4)
print(report2)
print("Macro average Test Precision, Recall and F1-Score...")
report3=metrics.precision_recall_fscore_support(best_labels, preds, average='macro')
print(report3)
print("Micro average Test Precision, Recall and F1-Score...")
report4=metrics.precision_recall_fscore_support(best_labels, preds, average='micro')
print(report4)
print("Weighted average Test Precision, Recall and F1-Score...")
report5=metrics.precision_recall_fscore_support(best_labels, preds, average='weighted')
print(report5)
print("Kappa")
Kappa = metrics.cohen_kappa_score(best_labels, preds)
print(Kappa)

